"""New proceeding; refuses incomplete or unrefined study summaries."""
import ast,json,importlib.util
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from model import ROOT,PROJECT
import manuscript_content as text

OUT=ROOT/'output';FIG=OUT/'figures';body_text=[]
AUTHORS='Joel Whaling, Aarav Patel*, Mandy Hung*, Yuancheng Fang*, Leona Cai*, Zenan Li*, Advik Raja*, Nathan Xu*, Linda Shi'
AFFILIATION='Institute of Engineering in Medicine, University of California San Diego, USA'


def max_direction(rows,key):
    return max((r for r in rows if r['angle_deg'] is not None),key=lambda r:r[key])


def refs():
    spec=importlib.util.spec_from_file_location('earlier_reference_text',PROJECT/'manuscript_content.py')
    old=importlib.util.module_from_spec(spec);spec.loader.exec_module(old)
    metadata={r['role']:r for r in json.loads((ROOT/'sources/reference_metadata.json').read_text(encoding='utf-8'))}
    metadata['epidermal']=json.loads((ROOT/'sources/epidermal_electronics_crossref.json').read_text(encoding='utf-8'))['metadata']
    def citation(role):
        a=metadata[role]
        authors=[]
        for name in a['author']:
            initials=' '.join('-'.join(part[0]+'.' for part in word.split('-') if part)
                              for word in name.get('given','').split() if word)
            authors.append(initials+' '+name.get('family',''))
        author=authors[0] if len(authors)==1 else (' and '.join(authors) if len(authors)==2 else ', '.join(authors[:-1])+', and '+authors[-1])
        year=a['published']['date-parts'][0][0]
        page=a.get('page')
        locator='article '+a['article-number'] if a.get('article-number') else (
            ('pp. ' if '-' in page else 'article ' if page.startswith('0') else 'p. ')+page)
        issue=', no. '+a['issue'] if a.get('issue') else ''
        return f'{author}, "{a["title"][0]}," {a["container-title"][0]}, vol. {a["volume"]}{issue}, {locator}, {year}, doi: {a["DOI"]}.'
    return [citation('epidermal'),old.REFERENCES[2],citation('soft_interlayer'),citation('motion_artifacts'),
            citation('PDMS'),citation('skin_anisotropy'),
            'Goodfellow, "Gold pellets," bulk material-property data. [Online]. Available: https://www.goodfellow.com/usa/gold-pellets-group. Accessed: Oct. 18, 2026.',
            'C. Geuzaine and J.-F. Remacle, "Gmsh: A 3-D finite element mesh generator with built-in pre- and post-processing facilities," International Journal for Numerical Methods in Engineering, vol. 79, pp. 1309-1331, 2009, doi: 10.1002/nme.2579.',
            'G. Dhondt, CalculiX: A Three-Dimensional Structural Finite Element Program, version 2.17 (Ubuntu package 2.17-3). [Online]. Available: https://www.dhondt.de/. Accessed: Oct. 18, 2026.',
            'T. Gustafsson and G. D. McBain, "scikit-fem: A Python package for finite element assembly," Journal of Open Source Software, vol. 5, article 2369, 2020, doi: 10.21105/joss.02369.']


def main():
    global doc
    result=json.loads((OUT/'Final_Study_Summary.json').read_text())
    assert result['material_configurations']==27 and result['directional_responses']==972
    assert len(result['mesh_comparisons'])>=4
    assert result['maximum_free_DOF_force_residual']<1e-5 and result['maximum_relative_energy_error']<5e-4
    required_figures=['Contact_Geometry','Assembly_On_Skin','Directional_Wire_Strain','Mechanical_Tradeoffs']
    assert all((FIG/(name+'.png')).exists() for name in required_figures)
    doc=Document(PROJECT/'sources/Avestia_Template.docx');first=deepcopy(doc.sections[0]._sectPr)
    for child in list(doc._element.body):
        if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        s=doc.styles[name];s.font.name='Cambria';s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(11)
        s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(0);s.paragraph_format.line_spacing=1
        s.paragraph_format.page_break_before=False
    doc.styles['Normal'].paragraph_format.first_line_indent=Cm(1)
    doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.bold=True
    for name in ['Heading 1','Heading 2']:
        s=doc.styles[name];s.font.bold=True;s.font.size=Pt(12 if name=='Heading 1' else 11)
        s.font.all_caps=False;s.font.small_caps=False;s.paragraph_format.first_line_indent=Pt(0)
        s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(3);s.paragraph_format.keep_with_next=True
    # Reuse the retained template's established paragraph/table helpers, not its old study.
    parsed=ast.parse((PROJECT/'build_paper.py').read_text(encoding='utf-8'))
    selected=[n for n in parsed.body if isinstance(n,ast.FunctionDef) and n.name in ['para','paragraphs','h','figure','table']]
    exec(compile(ast.Module(body=selected,type_ignores=[]),'template_helpers','exec'),globals())
    pa=para(text.TITLE,22,bold=True,center=True,indent=False,style='Title');pa.paragraph_format.space_after=Pt(12)
    pa=para(AUTHORS,12,bold=True,center=True,indent=False);pa.paragraph_format.space_after=Pt(5)
    para(AFFILIATION,11,center=True,indent=False)
    para('9500 Gilman Drive, La Jolla, CA 92093, USA',11,center=True,indent=False)
    para('*High school students participating in the IEM OPALS program',10,center=True,indent=False)
    pa=doc.add_paragraph();pa.paragraph_format.space_after=Pt(8);pa._p.get_or_add_pPr().append(first)
    nfront=sum(o['sampled_Pareto_nondominated'] for o in result['options'])
    change=result['maximum_mesh_changes_percent']
    abstract=(f'Abstract - Mechanical stability of a wearable biosensor depends on its component load paths, attachment, and imposed skin motion. We retained a sixteen-body CAD assembly and compared uniform, single-compliant, and graded supports at fixed thickness. The expanded three-dimensional linear-elastic study varies adhesive modulus, effective skin modulus, and extension direction. A geometric audit distinguishes the supplied contact from an explicitly gap-filled adhesive scenario. CalculiX/SPOOLES solves, independent high-contrast benchmarks, local-equilibrium checks, and direct mixed-load comparisons support numerical acceptance. The main filled-contact matrix contains 27 material configurations and 972 sampled directional responses reconstructed from tensor basis solutions; these are not 972 independent solves. Wire principal-strain demand, disturbance relative to matched bare-skin finite-element references, and adhesive distortion quantify separate tradeoffs. The sampled worst-case comparison identifies {nfront} nondominated material/support options. Across selected nominal and sampled-worst configurations, the largest final-refinement changes are {change["wire"]:.1f}%, {change["skin"]:.1f}%, and {change["adhesive"]:.1f}% for the three metrics. Results guide conditional prototype comparisons rather than establishing a universal graded-layer advantage. Unidentified component materials remain surrogates, bonding is idealized, and no electrical performance, clinical benefit, or ANSYS-equivalent physical accuracy is validated.')
    para(abstract,10,italic=True,indent=False)
    para('Keywords: Wearable biosensors, three-dimensional finite elements, compliant supports, skin attachment, material uncertainty.',10,indent=False)
    h('1. Introduction');paragraphs(text.INTRODUCTION)
    h('2. Methods');h('2.1. Assembly and attachment scenarios',True);paragraphs(text.GEOMETRY)
    figure('Contact_Geometry','Figure 1. Actual shared skin–adhesive facets in the medium meshes, accumulated in 1 mm² projected bins. Colors show geometric surface-area density, not pressure or adhesion strength. White bins contain no shared facets. Nonplanar surface area can exceed projected area. The supplied and filled scenarios are kept separate.')
    h('2.2. Material and motion variables',True);paragraphs(text.MATERIALS)
    table('Table 1. Isotropic inputs. Unidentified fixed properties are surrogates; no manufactured grade is established.',['Region','E (MPa)','ν'],[
        ['Skin','0.10, 0.20, 0.50','0.49'],['Adhesive','0.05, 0.10, 0.50','0.49'],['PDMS / cover','1.80','0.49'],['Soft region','0.10','0.49'],['Medium region','0.50','0.49'],['ABS','2,300','0.35'],['Sensor packages','10,000','0.30'],['Transmitter','20,000','0.30'],['Gold-labeled solids','78,500 [7]','0.42']])
    paragraphs(text.LOADING)
    for expression in ['c₁ = cos²θ − ν sin²θ','c₂ = sin²θ − ν cos²θ','c₃ = 2(1 + ν) sinθ cosθ']:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
        math=OxmlElement('m:oMath');run=OxmlElement('m:r');t=OxmlElement('m:t');t.text=expression;run.append(t);math.append(run);p._p.append(math)
    h('2.3. Discretization and numerical verification',True);paragraphs(text.VERIFICATION)
    h('2.4. Mechanical response measures',True);paragraphs(text.METRICS)
    figure('Assembly_On_Skin','Figure 2. Refined filled-contact assembly with uniform support at nominal material inputs: (a) undeformed geometry; (b) the sampled direction maximizing its wire strain percentile. The cover is hidden. Deformation is enlarged 10 times, while stress uses the actual 1% prescribed extension. Electronic-body colors show element-mean integration-point von Mises stress, saturated above 5 MPa; skin and supports retain material colors.')
    h('3. Results');h('3.1. Verification and refinement',True)
    independent=json.loads((ROOT/'verification/independent_high_contrast.json').read_text())
    manufactured=json.loads((ROOT/'verification/residual_refinement_verification.json').read_text())
    para(f'The accepted study contains {result["accepted_assembly_solves"]} distinct assembly load cases and {result["auxiliary_residual_correction_solves"]} auxiliary defect-correction solves. The latter improve the solution of existing cases rather than adding physical loading scenarios. Independent high-contrast displacement comparisons differ by at most {max(a["independent_displacement_relative_error"] for a in independent):.2g} in relative norm. The manufactured-error test recovers all checked fields to within {max(manufactured["recovery_relative_errors"].values()):.2g} relative error. Across accepted assembly cases, the maximum free-DOF relative force residual is {result["maximum_free_DOF_force_residual"]:.2g}, and the largest constrained-work energy discrepancy is {result["maximum_relative_energy_error"]:.2g}. These numerical checks do not calibrate biological or manufactured materials.')
    para('Table 2 compares sampled directional maxima on the medium and refined meshes for nominal architectures, the supplied-contact control, and configurations selected from the medium grid for worst responses. These checks sample the uncertainty space; they do not refine every grid point or establish certified error bounds. A small apparent design advantage must be interpreted alongside the measured mesh dependence.')
    para(f'Independent element-mean stress and strain comparisons have a maximum relative discrepancy of {result["maximum_independent_element_mean_tensor_error"]:.2g}, measured with volume-weighted Frobenius norms. Element means avoid assuming identical quadrature-point numbering between the two implementations; this is not a pointwise error bound for the biosensor assembly.')
    para(f'Each direct mixed-load assembly comparison must satisfy a relative field-norm discrepancy below 2e-5 for displacement, internal force, stress, and strain. The largest observed discrepancy across these fields is {max(result["maximum_direct_superposition_relative_errors"].values()):.2g}. Assembly free-DOF residuals must remain below 1e-5 and relative constrained-work energy errors below 5e-4; a completed solver execution alone does not meet these gates.')
    table('Table 2. Relative changes (%) in sampled directional maxima after refinement. F/S denote filled/supplied contact; U/C/G denote uniform/single-compliant/graded. Moduli are skin/adhesive in MPa.',['Case','Wire','Skin','Adhesive'],[
        [f'{"F" if a["geometry"]=="Filled" else "S"}/{dict(Uniform="U",Single="C",Graded="G")[a["design"]]} {a["skin_E_MPa"]:g}/{a["adhesive_E_MPa"]:g}',*[f'{a[k]["relative_change_percent"]:.1f}' for k in ['wire','skin','adhesive']]] for a in result['mesh_comparisons']])
    h('3.2. Directional and material response',True)
    ranges=[]
    for design,rows in result['filled_contact_nominal'].items():
        values=[a['wire_p95_absolute_principal_strain'] for a in rows if a['angle_deg'] is not None]
        ranges.append(f'{design.lower()}: {100*min(values):.3g}–{100*max(values):.3g}%')
    para('At nominal skin and adhesive moduli in the filled-contact geometry, the sampled wire strain percentiles span '+ '; '.join(ranges)+'. Figure 3 shows the directional response at nominal adhesive modulus, with bands spanning the selected skin stiffnesses. The bands represent parameter dependence, not statistical confidence intervals or a patient population.')
    figure('Directional_Wire_Strain','Figure 3. Wire absolute principal-strain 95th percentile versus extension direction for filled contact and adhesive modulus 0.10 MPa. Lines use skin modulus 0.20 MPa; shading encloses the three sampled skin-modulus curves at 0.10, 0.20, and 0.50 MPa. The end point repeats the equivalent 0° direction. All curves use 1% prescribed principal extension.')
    extra_ranges=[]
    for label,description in [('Equibiaxial_0.5percent','0.5% equal biaxial extension'),
                              ('Engineering_shear_1percent','1% engineering shear')]:
        values=[next(r for r in rows if r['loading']==label)['wire_p95_absolute_principal_strain']
                for rows in result['filled_contact_nominal'].values()]
        extra_ranges.append(f'{description}: {100*min(values):.3g}–{100*max(values):.3g}%')
    para('For nominal materials, the additional loadings give wire strain percentiles across the three support architectures of '+ '; '.join(extra_ranges)+'. Both additional tensors have maximum absolute prescribed in-plane principal strain 0.5%, versus 1% in the directional screen. They are separate motion comparisons and do not enter its directional envelopes or worst-case Pareto coordinates.')
    h('3.3. Attachment and mechanical tradeoffs',True)
    for design in ['Uniform','Single','Graded']:
        key='wire_p95_absolute_principal_strain'
        a=max_direction(result['supplied_contact_nominal'][design],key)[key]
        b=max_direction(result['filled_contact_nominal'][design],key)[key]
        para(f'For {design.lower()} support at nominal materials, filling the specified geometric clearance changes the sampled maximum wire strain percentile from {100*a:.3g}% to {100*b:.3g}% ({100*(b/a-1):+.1f}%). This comparison includes the explicitly added adhesive and changed shared surface; it is not a measured adhesion-strength effect.')
    table('Table 3. Sampled worst responses across three skin stiffnesses and 36 extension directions for each filled-contact design option. Strains are percentages; skin D is RMS perturbation divided by 0.01.',['Support','Adh. E','Wire ε95','Skin D','Adh. εdev'],[
        [a['design'],f'{a["adhesive_E_MPa"]:g}',f'{100*a["worst_wire"]:.3g}',f'{a["worst_skin"]:.3g}',f'{100*a["worst_adhesive"]:.3g}'] for a in result['options']])
    options=[f'{a["design"].lower()} / {a["adhesive_E_MPa"]:g} MPa adhesive' for a in result['options'] if a['sampled_Pareto_nondominated']]
    para('The sampled worst-case Pareto set contains '+ '; '.join(options)+'. Each objective is maximized separately over the selected skin stiffnesses and directions, so the three worst coordinates need not occur in the same loading state. Nondominance expresses a tradeoff within these metrics; it does not identify a universally superior device or assign clinical importance to the objectives.')
    figure('Mechanical_Tradeoffs','Figure 4. Filled-contact material/support options compared by sampled worst wire strain, skin disturbance, and adhesive RMS deviatoric strain. Color and marker fill identify support architecture; marker shape identifies adhesive modulus. Black outlines identify sampled nondominated options. Each coordinate is its own worst sampled value; maxima need not co-occur. Close apparent dominance must be considered with mesh sensitivity.')
    h('4. Discussion')
    para('The practical design question is how to reduce conductor deformation while preserving an acceptable mechanical interaction with skin and adhesive. The expanded comparison makes those responses explicit instead of replacing them with a whole-assembly stress average. A support or adhesive choice can move the design favorably in one coordinate while worsening another. Placement relative to motion is likewise a controlled variable whose effect should be measured before attributing improved performance solely to a graded layer.')
    para('The attachment audit also changes the interpretation of the supplied geometry. Its projected adhesive footprint and its shared bonded surface are different quantities. The filled scenario tests a specified alternative attachment construction without moving the electronics or altering support thickness. The resulting differences indicate what should be checked in manufactured specimens: actual bonding geometry, constituent properties, conductor deformation, and skin motion. They do not determine whether either construction will remain attached during use.')
    para('For prototype development, the sampled tradeoff set provides candidates for controlled comparison rather than a prescribed clinical design. Bench tests can measure full-field skin and conductor deformation at several mounting directions and compare adhesive formulations at matched geometry. Electrical measurements during those tests are necessary to connect mechanical demand to signal stability. The uncertainty grid and refinement results indicate where a favorable nominal response needs further confirmation.')
    h('5. Limitations');paragraphs(text.LIMITATIONS)
    h('6. Conclusion')
    para(f'This study extends the same sixteen-body multilayer biosensor assembly to directional motion, selected skin and adhesive stiffnesses, and explicit attachment scenarios. Verified tensor-basis solutions support 972 sampled directional responses in the main material matrix. Wire strain, disturbance relative to matched bare skin, and adhesive distortion identify {nfront} sampled nondominated options and make component/interface tradeoffs reviewable. The result is a conditional mechanical screening framework for prototype selection, with retained inputs and numerical checks. It does not establish biomarker sensing accuracy, comfort, fatigue life, or clinical benefit.')
    h('Data and code availability')
    para('The accompanying research package retains the original CAD, specified gap-filled construction, source scripts, solver inputs and execution records, verification cases, response tables, refined comparisons, and selected field arrays. Rejected and corrected numerical outputs are distinguished. All results are computational; no human or animal experiments were performed in this study.')
    h('References')
    references=refs()
    for i,s in enumerate(references,1):
        p=para(f'[{i}] {s}',10,indent=False);p.paragraph_format.left_indent=Cm(.55);p.paragraph_format.first_line_indent=Cm(-.55);p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for section in doc.sections:
        section.different_first_page_header_footer=False
        for hf in [section.header,section.first_page_header,section.even_page_header,section.footer,section.first_page_footer,section.even_page_footer]:
            for child in list(hf._element):hf._element.remove(child)
            hf._element.append(OxmlElement('w:p'))
        p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    doc.core_properties.title=text.TITLE;doc.core_properties.author=AUTHORS.replace('*','');doc.core_properties.comments=''
    doc.save(OUT/'Biosensor_Advanced_Proceeding.docx')
    (OUT/'Advanced_Manuscript_Text.txt').write_text('\n\n'.join(body_text),encoding='utf-8')
    stats=dict(words=len(' '.join(body_text).split()),figures=4,tables=3,references=len(references),accepted_assembly_load_cases=result['accepted_assembly_solves'])
    (OUT/'Manuscript_Statistics.json').write_text(json.dumps(stats,indent=2));print(json.dumps(stats))


if __name__=='__main__':
    main()
