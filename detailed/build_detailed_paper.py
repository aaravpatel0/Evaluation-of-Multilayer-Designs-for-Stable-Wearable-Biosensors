from pathlib import Path
from copy import deepcopy
import sys,ast,json
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.collections import PolyCollection
from matplotlib.colors import Normalize
from matplotlib.cm import ScalarMappable
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from study_analysis import get,part,group,audit
from run_ccx import ROOT,NAMES
from skin_stress_figure import main as build_skin_stress_figure
import paper_content as text
sys.path.insert(0,str(ROOT.parent));import manuscript_content as old
OUT=ROOT/'output';OUT.mkdir(exist_ok=True);FIG=OUT/'figures';FIG.mkdir(exist_ok=True)
build_skin_stress_figure()
names=['Uniform','Single','Graded'];colors=['#374b67','#157b8c','#b55f30'];short={4:'ABS body',5:'Sensor 1',6:'Sensor 2',7:'Sensor 3',8:'Sensor 4',9:'Gold feature 1',10:'Transmitter',11:'Gold feature 2',12:'ABS feature',13:'Gold wire 1',14:'Gold wire 2',15:'Sensor 5'}
report=audit();base=[get(d) for d in names];medium=[get(d,h=1.7) for d in names]
wire=[group(a,[13,14]) for a in base];mw=[group(a,[13,14]) for a in medium];increases=[100*(v/wire[0]-1) for v in wire]
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False,'axes.labelsize':9,'legend.fontsize':8,'xtick.labelsize':8,'ytick.labelsize':8})
def save(fig,name):fig.savefig(FIG/(name+'.png'),dpi=330,bbox_inches='tight');fig.savefig(FIG/(name+'.pdf'),bbox_inches='tight');plt.close(fig)

fig,axs=plt.subplots(2,1,figsize=(3.6,3.2),sharex=True)
for name,col in zip(names,colors):
    vals=[get(name,h=h) for h in [2.5,1.7,1.2]];xx=[v['elements']/1000 for v in vals];axs[0].plot(xx,[group(v,[13,14]) for v in vals],'o-',color=col,label=name,ms=3);axs[1].plot(xx,[max(part(v,13)['max_gp_vm_MPa'],part(v,14)['max_gp_vm_MPa']) for v in vals],'o-',color=col,ms=3)
axs[0].set_ylabel('Mean stress (MPa)');axs[1].set_ylabel('Max. stress (MPa)');axs[1].set_xlabel('Total quadratic tetrahedra (thousands)');axs[0].legend(frameon=False);fig.tight_layout(h_pad=.8);save(fig,'Mesh_Convergence')

fig,ax=plt.subplots(figsize=(3.6,2.6));xx=np.arange(3);modes=['Tension','Compression','Bending']
for i,(name,col) in enumerate(zip(names,colors)):ax.bar(xx+(i-1)*.24,[group(get(name,mode=m,h=1.7),[13,14]) for m in modes],.24,color=col,label=name)
ax.set_xticks(xx,['Extension','Shortening','Transverse']);ax.set_ylabel('Wire mean stress (MPa)');ax.legend(frameon=False);fig.tight_layout();save(fig,'Loading_Comparison')

mesh=np.load(ROOT/'meshes/h1.2/mesh.npz');p,c,r=mesh['points'],mesh['cells'],mesh['region'];ids=np.where(r==13)[0];cw=c[ids,:4]
allfaces=np.vstack([cw[:,[0,1,2]],cw[:,[0,1,3]],cw[:,[0,2,3]],cw[:,[1,2,3]]]);owners=np.tile(ids,4);unique,index,counts=np.unique(np.sort(allfaces,axis=1),axis=0,return_index=True,return_counts=True);ff=unique[counts==1];own=owners[index[counts==1]];tri=p[ff];order=np.argsort(tri[:,:,2].mean(axis=1));norm=Normalize(0,15)
fig,axs=plt.subplots(3,1,figsize=(3.6,3.65),sharex=True)
for ax,name,a in zip(axs,names,base):
    dest=ROOT/'runs'/f'h1.2_{name}_Tension_LE_P10000_A0.1';fields=np.load(dest/'fields.npz');eid=fields['element_id']-1;vm=fields['vm'];element_mean=np.bincount(eid,weights=vm,minlength=len(c))/np.maximum(np.bincount(eid,minlength=len(c)),1)
    coll=PolyCollection(tri[order,:,:2],array=element_mean[own[order]],norm=norm,cmap='magma',edgecolors='none');ax.add_collection(coll);ax.set_xlim(6.9,12.1);ax.set_ylim(-6.2,-3.3);ax.set_aspect('equal');ax.set_ylabel('y (mm)');ax.set_title(name,fontsize=9,loc='left',pad=1)
axs[-1].set_xlabel('x (mm)');fig.subplots_adjust(left=.16,right=.83,hspace=.55,top=.94,bottom=.14);cb=fig.colorbar(ScalarMappable(norm=norm,cmap='magma'),cax=fig.add_axes([.88,.18,.035,.68]),extend='max');cb.set_label('Equivalent stress (MPa)');save(fig,'Wire_Stress_Field')

template=ROOT.parent/'sources/Avestia_Template.docx';doc=Document(template);first=deepcopy(doc.sections[0]._sectPr)
for child in list(doc._element.body):
    if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Cambria';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(11);st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=1
doc.styles['Normal'].paragraph_format.first_line_indent=Cm(1)
for st in ['Heading 1','Heading 2']:
    s=doc.styles[st];s.font.bold=True;s.font.size=Pt(12 if st=='Heading 1' else 11);s.font.all_caps=False;s.font.small_caps=False;s.paragraph_format.first_line_indent=Pt(0);s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(3);s.paragraph_format.keep_with_next=True
doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.bold=True;body_text=[]
# Reuse the previously visually verified Word-formatting helpers without executing its 2D study.
source=ast.parse((ROOT.parent/'build_paper.py').read_text());helpers=[n for n in source.body if isinstance(n,ast.FunctionDef) and n.name in ['para','paragraphs','h','figure','table']];exec(compile(ast.Module(body=helpers,type_ignores=[]),'formatting_helpers','exec'),globals())

pa=para(text.TITLE,22,bold=True,center=True,indent=False,style='Title');pa.paragraph_format.space_after=Pt(12)
pa=para(old.AUTHORS,12,bold=True,center=True,indent=False);pa.paragraph_format.space_after=Pt(5)
para(old.AFFILIATION,11,center=True,indent=False);para('9500 Gilman Drive, La Jolla, CA 92093, USA',11,center=True,indent=False);para(old.OPALS,10,center=True,indent=False)
pa=doc.add_paragraph();pa.paragraph_format.space_after=Pt(8);pa._p.get_or_add_pPr().append(first)
meshchange=max(100*abs(a/b-1) for a,b in zip(wire,mw))
abstract=f'''Abstract - The mechanical performance of a wearable biosensor depends on the load paths through its individual components and interconnects. We evaluated a sixteen-body STEP assembly containing skin, adhesive, PDMS, ABS bodies, five sensor components, a transmitter, patterned gold features, two routed gold wires, and a protective cover. Three equal-thickness support architectures were compared using Gmsh meshes and CalculiX 2.23 structural simulations, with Python post-processing. Unidentified component properties were assigned as explicit mechanical surrogates and varied in sensitivity cases. Quadratic tetrahedral meshes reached {report['finest_mesh_elements']:,} elements and {report['finest_mesh_dofs']:,} displacement degrees of freedom. Analytical cube benchmarks and an independent scikit-fem implementation verified basic solver consistency. At 1% prescribed skin extension in the linear model, combined wire mean von Mises stresses were {wire[0]:.3f}, {wire[1]:.3f}, and {wire[2]:.3f} MPa for uniform, single-buffer, and graded supports. The buffered cases increased this wire metric by {increases[1]:.1f}% and {increases[2]:.1f}% relative to uniform PDMS. Wire means changed by up to {meshchange:.1f}% between the last two meshes, while local maxima were more sensitive. Stiffness perturbations assessed material dependence; exploratory nonlinear attempts did not establish full-load equilibrium. The supplied ABS stem crosses the support and partly enters the adhesive, providing a distinct local load path. These results do not establish a general benefit from compliant grading. They support component-specific assessment with explicit geometric, material, and numerical assumptions. No electrical function, fatigue life, or equivalence to an ANSYS model was validated.'''
para(abstract,10,italic=True,indent=False);para('Keywords: Wearable biosensors, three-dimensional finite elements, gold interconnects, compliant supports, verification.',10,indent=False)
h('1. Introduction');paragraphs(text.INTRO)
h('2. Methods');h('2.1. Assembly and controlled support variants',True);paragraphs(text.METHOD_GEOMETRY)
figure('Sensor_On_Skin_3D','Figure 1. Three-dimensional views of the accepted finest mesh: (a) assembly on the modeled skin slab; (b) uniform-support case at 1% prescribed skin extension. Displacements are enlarged 10 times; no nonlinear equilibrium is implied. Only electronic bodies carry the stress colors, using adjacent-element mean integration-point von Mises values. Values above 5 MPa are saturated. Skin and supports retain material colors; the top cover is hidden in both views.')
h('2.2. Material definitions',True);paragraphs(text.METHOD_MATERIALS)
table('Table 1. Representative isotropic inputs. Unidentified package and transmitter properties are surrogate assumptions.',['Region','E (MPa)','Poisson ratio'],[['Skin','0.20','0.49'],['Adhesive','0.10','0.49'],['PDMS / cover','1.80','0.49'],['Soft region','0.10','0.49'],['Medium region','0.50','0.49'],['ABS','2,300','0.35'],['Sensor packages','10,000','0.30'],['Transmitter','20,000','0.30'],['Gold','78,500','0.42']])
h('2.3. Software and loading',True);paragraphs(text.METHOD_SOLVER)
h('2.4. Discretization and verification',True);paragraphs(text.METHOD_NUMERICS)
mesha=[json.loads((ROOT/'meshes'/f'h{v:g}'/'mesh_audit.json').read_text()) for v in [2.5,1.7,1.2]]
table('Table 2. Quadratic tetrahedral mesh sequence. Dimensions are in mm.',['Global size','Wire size','Elements','DOFs'],[[f"{a['h_mm']:.1f}",f"{a['wire_size_mm']:.2f}",f"{a['elements']:,}",f"{a['dofs']:,}"] for a in mesha])
h('3. Results');h('3.1. Numerical verification and refinement',True)
verification=json.loads((ROOT/'verification/verification_summary.json').read_text());berr=max(a['independent_displacement_relative_error'] for a in verification)
para(f'''The homogeneous benchmark reactions agreed with the analytical value of 0.010000 N to within {max(a.get('analytical_reaction_relative_error',0) for a in verification):.1g} relative error. The independent CalculiX and scikit-fem displacement vectors differed by at most {berr:.2g} in relative Euclidean norm across the homogeneous and heterogeneous benchmark cases. These comparisons use printed solver output and are limited by its precision. They verify the selected benchmark problems, not a physical device.''')
para(f'''All accepted assembly meshes had positive Jacobians. The finest mesh contained one connected finite-element assembly. Across the completed assembly cases, the largest normalized total-force imbalance was {report['maximum_relative_force_balance']:.2g}. Figure 2 shows the wire mesh histories. The combined mean wire stress changed by {100*abs(wire[0]/mw[0]-1):.1f}%, {100*abs(wire[1]/mw[1]-1):.1f}%, and {100*abs(wire[2]/mw[2]-1):.1f}% from the intermediate to finest mesh for uniform, single-buffer, and graded supports.''')
figure('Mesh_Convergence','Figure 2. Linear-extension mesh histories. Wire means combine both wire volumes; maxima are the largest wire integration-point values. Local wire resolution is refined together with global mesh size.')
para('The minimum Gmsh element-quality statistic, minSICN, was 0.0074, 0.0156, and 0.0263 across the three meshes. Positive Jacobians do not eliminate element-shape error. Successive-mesh differences measure numerical sensitivity, not certified error bounds. Component-level changes, including upper-tail metrics, are retained in the audit. Wire maxima and junction stresses require particular caution because their convergence can differ from that of volume means.')
h('3.2. Component stresses under extension',True)
para(f'''Table 3 reports the finest-mesh mean equivalent stress for every electronic body. The two-wire mean was {wire[0]:.3f} MPa with uniform PDMS, {wire[1]:.3f} MPa with the single compliant region, and {wire[2]:.3f} MPa with grading. Thus the nominal buffered designs increased this measure by {increases[1]:.1f}% and {increases[2]:.1f}%, respectively. A ranking based on a homogeneous sensing layer would not characterize this routed-wire response.''')
table('Table 3. Finest-mesh component mean von Mises stress in linear extension (MPa).',['Component','Uniform','Single','Graded'],[[short[r],*[f"{part(a,r)['mean_vm_MPa']:.3f}" for a in base]] for r in range(4,16)])
para(f'''The single-to-graded wire difference was {100*abs(wire[2]/wire[1]-1):.1f}%, smaller than the measured changes on the last mesh refinement. That small separation is insufficient by itself to establish a robust advantage for grading. Some sensor-package means decreased with compliant supports even as wire means increased, so the assembly exhibits a component-level tradeoff. Several package means changed more than wire means during refinement; Table 3 should not be interpreted as a uniformly converged stress field.''')
figure('Wire_Stress_Field','Figure 3. Plan views of gold wire 1 under extension, colored by the adjacent element mean integration-point stress. The common scale spans 0-15 MPa; higher values are saturated. No cross-material nodal smoothing is used.')
para(f'''Maximum wire integration-point stresses were {max(part(base[0],13)['max_gp_vm_MPa'],part(base[0],14)['max_gp_vm_MPa']):.2f}, {max(part(base[1],13)['max_gp_vm_MPa'],part(base[1],14)['max_gp_vm_MPa']):.2f}, and {max(part(base[2],13)['max_gp_vm_MPa'],part(base[2],14)['max_gp_vm_MPa']):.2f} MPa. Figure 3 resolves spatial variation along the routed conductor, including its turns and attachment regions. These maxima are reported as numerical diagnostics and are not compared with an assumed universal yield or failure threshold.''')
para('Mean and maximum principal-strain measures for both wires and all sensor packages accompany the stress data. A small strain in a stiff package does not imply zero displacement or a motion-free package, since rigid translation and rotation do not produce the same response as internal deformation.')
h('3.3. Loading and nonlinear checks',True)
para(f'''Figure 4 compares mean wire stress in the three prescribed skin-edge fixtures on the intermediate mesh. The exact linear shortening solution is proportional to extension with reversed sign and half the amplitude. Calculated component mean equivalent stresses followed this halving relation with a maximum discrepancy of {100*report['compression_scaling_relative_error']:.3f}%. This is a numerical consistency observation, not independent evidence about compression stability.''')
figure('Loading_Comparison','Figure 4. Mean wire stress for the intermediate-mesh fixtures. Extension, shortening, and transverse motion use the prescribed displacements stated in Methods. The transverse case is not pure bending.')
para('Exploratory finite-geometry attempts on the uniform-support intermediate mesh showed growing equilibrium residuals from both zero and linear-field initializations. The attempts were halted without establishing full-load nonlinear equilibrium; their stress fields are excluded from all tables and figures. This outcome does not verify the linear approximation at 1% nominal extension. It identifies a need for further load continuation, constitutive validation, and independent nonlinear reconstruction before finite-deformation predictions can be made.')
h('3.4. Unidentified-material sensitivity',True)
sens=[('Nominal',get(h=1.7)),('Package 2 GPa',get(h=1.7,package=2000.)),('Package 130 GPa',get(h=1.7,package=130000.)),('Adhesive 0.05 MPa',get(h=1.7,adhesive=.05)),('Adhesive 0.50 MPa',get(h=1.7,adhesive=.5))]
table('Table 4. Uniform-support stiffness sensitivity in linear extension on the intermediate mesh. Mean stresses are in MPa.',['Perturbation','Wires','Sensor packages'],[[s,f'{group(a,[13,14]):.3f}',f'{group(a,[5,6,7,8,15]):.3f}'] for s,a in sens])
para('The sensitivity cases alter one assumed input at a time. They quantify model dependence within the selected ranges, while preserving the original CAD. They are not a distribution of manufactured devices or a replacement for identifying the actual package and adhesive materials.')
adh_vals=[group(get(h=1.7,adhesive=v),[13,14]) for v in [.05,.1,.5]];pkg_vals=[group(get(h=1.7,package=v),[5,6,7,8,15]) for v in [2000.,10000.,130000.]]
para(f'''Across the adhesive sweep, combined wire means ranged from {min(adh_vals):.3f} to {max(adh_vals):.3f} MPa. The package-stiffness sweep produced combined sensor-package means from {min(pkg_vals):.3f} to {max(pkg_vals):.3f} MPa. These ranges show why the assigned nominal properties must remain visible alongside the architecture comparison.''')
h('4. Discussion');paragraphs(text.DISCUSSION)
h('5. Limitations');paragraphs(text.LIMITS)
h('6. Conclusion')
para(f'''A sixteen-body three-dimensional biosensor assembly was meshed and simulated with explicit sensor bodies, transmitter geometry, gold features, and two routed wire solids. Under the specified linear 1% skin extension, the finest-mesh wire mean equivalent stress was {wire[0]:.3f} MPa for uniform PDMS and increased to {wire[1]:.3f} and {wire[2]:.3f} MPa for the single and graded compliant supports. The supplied assembly therefore does not support a general claim that adding compliance reduces interconnect stress. Results for each electronic body, rather than a whole-package average, are needed to assess the tradeoffs.''')
para('Analytical and independent numerical benchmarks support implementation consistency, and three mesh levels quantify response sensitivity. Material perturbations expose additional dependence on assumed properties. Exploratory nonlinear attempts did not establish full-load equilibrium, so the quantitative findings remain conditional on the linear formulation. These computations provide a traceable screening study, with no claim of validated electrical performance, component reliability, or ANSYS-equivalent predictive accuracy.')
h('Data and code availability');para('The accompanying package includes the original STEP geometry, material assignments, meshing and post-processing scripts, CalculiX input decks and logs, component tables, numerical audits, and selected full displacement and integration-point fields. An interactive mesh viewer permits inspection of the separate CAD bodies. All results are computational; no human or animal experiments were performed.')
references=[*old.REFERENCES[:5], 'Goodfellow, "Gold pellets," material-property data: tensile modulus 78.5 GPa and Poisson ratio 0.42. [Online]. Available: https://www.goodfellow.com/usa/gold-pellets-group. Accessed: Oct. 17, 2026.', 'C. Geuzaine and J.-F. Remacle, "Gmsh: A 3-D finite element mesh generator with built-in pre- and post-processing facilities," International Journal for Numerical Methods in Engineering, vol. 79, no. 11, pp. 1309-1331, 2009, doi: 10.1002/nme.2579.', 'G. Dhondt, CalculiX CrunchiX User\'s Manual, version 2.23. [Online]. Available: https://www.dhondt.de/ccx_2.23.pdf. Accessed: Oct. 17, 2026.', 'T. Gustafsson and G. D. McBain, "scikit-fem: A Python package for finite element assembly," Journal of Open Source Software, vol. 5, no. 52, article 2369, 2020, doi: 10.21105/joss.02369.']
references.insert(6,'INEOS Styrolution, "Styrenics product portfolio EMEA," 2026. [Online]. Available: https://www.ineos.com/contentassets/8ca5460f3c3049c3b7337f35ba8dc1ae/2026_styrenics_product_presentation_emea.pdf. Accessed: Oct. 17, 2026.')
h('References')
for i,s in enumerate(references,1):
    pp=para(f'[{i}] {s}',10,indent=False);pp.paragraph_format.left_indent=Cm(.55);pp.paragraph_format.first_line_indent=Cm(-.55);pp.paragraph_format.space_after=Pt(3);pp.alignment=WD_ALIGN_PARAGRAPH.LEFT
for section in doc.sections:
    section.different_first_page_header_footer=False
    for hf in [section.header,section.first_page_header,section.even_page_header,section.footer,section.first_page_footer,section.even_page_footer]:
        for child in list(hf._element):hf._element.remove(child)
        hf._element.append(OxmlElement('w:p'))
    pp=section.footer.paragraphs[0];pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');pp._p.append(fld)
doc.core_properties.title=text.TITLE;doc.core_properties.author=old.AUTHORS.replace('*','');doc.core_properties.comments='';doc.save(OUT/'Biosensor_Proceeding_3D.docx')
(OUT/'Manuscript_3D_Text.txt').write_text('\n\n'.join(body_text),encoding='utf-8');(OUT/'Paper_Statistics.json').write_text(json.dumps(dict(words=len(' '.join(body_text).split()),figures=4,tables=4,references=len(references),assembly_cases=report['solved_assembly_cases']),indent=2));print((OUT/'Paper_Statistics.json').read_text())
