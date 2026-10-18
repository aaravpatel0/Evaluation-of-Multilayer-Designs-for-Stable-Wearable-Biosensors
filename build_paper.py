from pathlib import Path
from copy import deepcopy
import csv,json,hashlib
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.collections import PolyCollection
from matplotlib.patches import Rectangle
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import manuscript_content as text

ROOT=Path(__file__).parent; OUT=ROOT/'output';OUT.mkdir(exist_ok=True);FIG=OUT/'figures';FIG.mkdir(exist_ok=True)
def read(name):
    rows=list(csv.DictReader((ROOT/'results'/name).open()))
    return [{k:(v if k in ['design','mode'] else float(v)) for k,v in row.items()} for row in rows]
rows=read('metrics.csv'); controls=read('controls.csv'); names=['Uniform','Single','Graded']; colors=['#374b67','#157b8c','#b55f30']
def get(design,mode='Tension',r=16,data=rows):
    return next(a for a in data if a['design']==design and a['mode']==mode and a['refinement']==r and a['soft_MPa']==.1 and a['skin_MPa']==.2 and a['nu']==.49)
base=[get(n) for n in names]; bend=[get(n,'Bending') for n in names]
cm=[get(n,data=controls) for n in ['Matched','Reverse']]
red=[100*(1-a['gold_mean_vm_MPa']/base[0]['gold_mean_vm_MPa']) for a in base]
verify=json.loads((ROOT/'results/verification.json').read_text())
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False,'axes.labelsize':9,'legend.fontsize':8,'xtick.labelsize':8,'ytick.labelsize':8})

def save(fig,name):
    fig.savefig(FIG/(name+'.png'),dpi=320,bbox_inches='tight');fig.savefig(FIG/(name+'.pdf'),bbox_inches='tight');plt.close(fig)

fig,axs=plt.subplots(3,1,figsize=(3.6,2.6),sharex=True)
for i,(ax,name) in enumerate(zip(axs,names)):
    ax.add_patch(Rectangle((0,0),5,.2,color='#e0c8b1'));ax.text(4.9,.1,'Skin upper surface',ha='right',va='center',fontsize=7)
    layers={'Uniform':[(.4,'#acbfd1','PDMS')],'Single':[(.2,'#a7d9d5','Soft'),(.2,'#acbfd1','PDMS')],'Graded':[(.1,'#a7d9d5','Soft'),(.1,'#efb68c','Medium'),(.2,'#acbfd1','PDMS')]}[name]
    y=.2
    for t,col,label in layers:
        ax.add_patch(Rectangle((0,y),5,t,color=col,ec='white',lw=.6))
        if t>=.2: ax.text(2.5,y+t/2,label,ha='center',va='center',fontsize=7)
        y+=t
    ax.add_patch(Rectangle((0,.6),1,.02,color='#b5901d'));ax.add_patch(Rectangle((1,.6),4,.02,color='#d6dde6'));ax.add_patch(Rectangle((0,.62),5,.1,color='#d6dde6'));ax.annotate('Gold',xy=(.5,.61),xytext=(.5,.84),ha='center',fontsize=7,arrowprops={'arrowstyle':'-','lw':.6})
    ax.text(2.5,.67,'Encapsulation',ha='center',va='center',fontsize=7);ax.set_title(name,loc='left',fontsize=9,pad=0);ax.set_ylim(0,.95);ax.set_xlim(0,5);ax.set_yticks([]);ax.spines['left'].set_visible(False)
axs[-1].set_xlabel('Position along skin x (mm)');fig.tight_layout(h_pad=.3);save(fig,'architecture')

fig,ax=plt.subplots(figsize=(3.6,2.4));xx=np.arange(3);w=.34
ax.bar(xx-w/2,[a['gold_mean_vm_MPa'] for a in base],w,label='Axial extension',color=colors)
ax.bar(xx+w/2,[a['gold_mean_vm_MPa'] for a in bend],w,label='Transverse case',color=colors,alpha=.4,hatch='//')
ax.set_xticks(xx,names);ax.set_ylabel('Gold mean equivalent stress (MPa)');ax.legend(frameon=False);ax.set_ylim(0,1.8);fig.tight_layout();save(fig,'mean_stress')

fig,axs=plt.subplots(2,1,figsize=(3.6,3.5),sharex=True)
for name,col in zip(names,colors):
    vals=[get(name,r=r) for r in [1,2,4,8,16]]
    axs[0].plot([50,100,200,400,800],[v['gold_mean_vm_MPa'] for v in vals],'o-',color=col,label=name,ms=3)
    axs[1].plot([50,100,200,400,800],[v['gold_max_cell_vm_MPa'] for v in vals],'o-',color=col,ms=3)
axs[0].set_ylabel('Mean stress (MPa)');axs[1].set_ylabel('Max. cell stress (MPa)');axs[1].set_xlabel('Elements along span');axs[1].set_xscale('log',base=2);axs[1].set_xticks([50,100,200,400,800],['50','100','200','400','800']);axs[0].legend(frameon=False);fig.tight_layout();save(fig,'mesh_study')

fig,axs=plt.subplots(3,1,figsize=(3.6,3.6),sharex=True)
for ax,name in zip(axs,names):
    dat=np.load(ROOT/'results'/f'{name}_Tension.npz');p=dat['points'];c=dat['cells'];gold=dat['gold'];v=dat['vm'];cc=c[gold]; polys=p[cc].copy();polys[:,:,1]-=2.4
    coll=PolyCollection(polys,array=v[gold],cmap='magma',edgecolors='none',clim=(0,4.6));ax.add_collection(coll);ax.set_xlim(0,1);ax.set_ylim(0,.02);ax.set_yticks([0,.02],['0','0.02']);ax.set_ylabel('y (mm)');ax.set_title(name,loc='left',pad=1,fontsize=9)
axs[-1].set_xlabel('Gold position x (mm)');fig.subplots_adjust(left=.18,right=.83,hspace=.6,bottom=.13,top=.95);cb=fig.colorbar(coll,cax=fig.add_axes([.87,.18,.035,.7]));cb.set_label('Equivalent stress (MPa)');save(fig,'gold_fields')

fig,ax=plt.subplots(figsize=(3.6,2.4)); ns=['Single','Graded','Matched','Reverse'];vals=[base[1],base[2],*cm];ax.bar(np.arange(4),[a['gold_mean_vm_MPa'] for a in vals],color=['#157b8c','#b55f30','#8b7eaa','#b7a880']);ax.set_xticks(np.arange(4),['Single','Graded','Matched','Reverse']);ax.set_ylabel('Gold mean stress in tension (MPa)');ax.set_ylim(0,.85);fig.tight_layout();save(fig,'controls')

template=ROOT/'sources/Avestia_Template.docx';doc=Document(template);first=deepcopy(doc.sections[0]._sectPr)
for child in list(doc._element.body):
    if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Cambria';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(11)
    st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0);st.paragraph_format.line_spacing=1
normal=doc.styles['Normal'];normal.paragraph_format.first_line_indent=Cm(1)
for st in ['Heading 1','Heading 2']:
    doc.styles[st].font.bold=True;doc.styles[st].font.size=Pt(12 if st=='Heading 1' else 11);doc.styles[st].paragraph_format.first_line_indent=Pt(0);doc.styles[st].paragraph_format.space_before=Pt(8);doc.styles[st].paragraph_format.space_after=Pt(3);doc.styles[st].paragraph_format.keep_with_next=True
    doc.styles[st].font.all_caps=False;doc.styles[st].font.small_caps=False
doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.bold=True
body_text=[]
def para(s='',size=11,italic=False,bold=False,center=False,indent=True,style=None):
    p=doc.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.first_line_indent=Cm(1) if indent else Pt(0)
    r=p.add_run(s);r.font.name='Cambria';r.font.size=Pt(size);r.font.italic=italic;r.font.bold=bold;body_text.append(s);return p
def paragraphs(a):
    for s in a:para(s)
def h(s,sub=False):
    p=doc.add_paragraph(s,style='Heading 2' if sub else 'Heading 1');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.left_indent=Pt(0);p.paragraph_format.first_line_indent=Pt(0)
    n=OxmlElement('w:numPr');nid=OxmlElement('w:numId');nid.set(qn('w:val'),'0');n.append(nid);p._p.get_or_add_pPr().append(n)
    for r in p.runs:r.font.all_caps=False;r.font.small_caps=False
    body_text.append(s)
def figure(name,caption):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(FIG/(name+'.png')),width=Inches(3.58));cp=para(caption,10,center=True,indent=False);cp.paragraph_format.keep_together=True
def table(caption,headers,data):
    p=para(caption,10,center=True,indent=False);p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for cell,title in zip(t.rows[0].cells,headers):cell.text=title
    for row in data:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for row_index,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.width=Inches(3.58/len(headers))
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=row_index<len(t.rows)-1;p.paragraph_format.keep_together=True
                for r in p.runs:r.font.name='Cambria';r.font.size=Pt(9)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell in t.rows[0].cells:
        for r in cell.paragraphs[0].runs:r.font.bold=True
        tcPr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');tcPr.append(shade)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)
def equation(s):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
    om=OxmlElement('m:oMath')
    def run(s):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);return r
    def sub(b,s):
        x=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(run(b));a=OxmlElement('m:sub');a.append(run(s));x.append(e);x.append(a);return x
    if s.endswith('(1)'):
        om.append(run('div σ = 0;   ε = (∇u + ∇uᵀ)/2;   '));om.append(sub('ε','zz'));om.append(run(' = 0     (1)'))
    elif s.endswith('(2)'):
        om.append(run('σ = 2μ'));om.append(sub('ε','dev'));om.append(run(' + κ tr(ε)I     (2)'))
    else:
        om.append(run('R = 100(1 − '));om.append(sub('S','d'));om.append(run(' / '));om.append(sub('S','0'));om.append(run(')     (3)'))
    p._p.append(om)

p=para(text.TITLE,22,bold=True,center=True,indent=False,style='Title');p.paragraph_format.space_after=Pt(12)
p=para(text.AUTHORS,12,bold=True,center=True,indent=False);p.paragraph_format.space_after=Pt(5)
para(text.AFFILIATION,11,center=True,indent=False)
para('9500 Gilman Drive, La Jolla, CA 92093, USA',11,center=True,indent=False)
para(text.OPALS,10,center=True,indent=False)
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8);p._p.get_or_add_pPr().append(first)
abstract=f'''Abstract - Compliant supports can reduce mechanical loading of stiff sensing elements in skin-mounted biosensors, but a graded buffer also changes overall support compliance. We compared uniform PDMS, a single soft buffer, and a two-layer graded buffer in a reproducible two-dimensional plane-strain finite-element study. Total support thickness was fixed at 0.40 mm, with identical skin, gold, and encapsulation geometry. Prescribed skin-edge motions represented axial extension, axial shortening, and transverse deformation. Selectively integrated quadrilateral elements were verified against analytical and affine benchmarks and refined from 50 to 800 elements along the span. Under 1% nominal extension, gold mean von Mises stress was {base[0]['gold_mean_vm_MPa']:.3f}, {base[1]['gold_mean_vm_MPa']:.3f}, and {base[2]['gold_mean_vm_MPa']:.3f} MPa for uniform, single-buffer, and graded supports, respectively. Corresponding stress reductions were {red[1]:.1f}% and {red[2]:.1f}%. The single buffer also produced the lowest mean stress in the prescribed transverse case. Mean stresses changed by less than 5% between the last two mesh levels, while localized maxima remained more sensitive. A homogeneous buffer matched to the graded pair's elementary shear compliance produced a similar mean sensor stress. The results support compliance-controlled mechanical isolation and do not establish a general advantage for grading. They provide a screening basis for subsequent nonlinear, interfacial, and experimental studies; electrical drift and failure were not evaluated.'''
para(abstract,10,italic=True,indent=False);para('Keywords: Wearable biosensors, compliant buffers, finite elements, mechanical isolation, mesh refinement.',10,indent=False)
h('1. Introduction');paragraphs(text.INTRO)
h('2. Related work');paragraphs(text.RELATED)
h('3. Methods');h('3.1. Geometry and material definitions',True);paragraphs(text.METHOD1)
figure('architecture','Figure 1. Support architectures with true layer thicknesses and skin cropped to its upper 0.20 mm. Colors: soft green, medium orange, PDMS blue, encapsulation gray, and gold yellow. Full skin thickness is 2.00 mm.')
table('Table 1. Nominal isotropic material inputs.',['Material','E (MPa)','Poisson ratio'],[['Skin','0.20','0.49'],['Soft buffer','0.10','0.49'],['Medium buffer','0.50','0.49'],['PDMS and cap','1.80','0.49'],['Gold','110,000','0.34']])
h('3.2. Formulation and boundary conditions',True);paragraphs(text.METHOD2)
equation('div σ = 0;   ε = (∇u + ∇uᵀ)/2;   εzz = 0     (1)')
equation('σ = 2μεdev + κ tr(ε)I     (2)')
para('Here u is displacement, ε is strain, σ is Cauchy stress in the linearized formulation, I is the identity tensor, mu = E/[2(1 + nu)], and kappa = E/[3(1 - 2nu)].')
h('3.3. Metrics and numerical study',True);paragraphs(text.METHOD3)
equation('R = 100(1 − ⟨σVM⟩design / ⟨σVM⟩uniform)     (3)')
para('Here R is percentage reduction, S with subscript d is the design gold-area-weighted mean equivalent stress, and S with subscript 0 is the uniform reference mean. A positive value indicates less stress than the reference at the same prescribed motion.')
h('4. Results');h('4.1. Verification and mesh sensitivity',True)
maxres=max(a['residual'] for a in rows+controls)
para(f'''The assembled affine benchmark reproduced displacement with a maximum absolute error of {verify['assembled_affine_displacement_error_mm']:.2g} mm. The analytical plane-strain tension benchmark gave a reaction of {verify['uniaxial_reaction_N']:.8f} N against {verify['analytical_reaction_N']:.8f} N analytically, a relative difference of {verify['reaction_relative_error']:.2g}. Rigid rotation produced zero strain to the precision of the test. Across the recorded model solves, the normalized free-equilibrium residual was at most {maxres:.2g}. These checks support implementation consistency; they do not substitute for geometric refinement.''')
para('Figure 2 separates the mesh history of the mean gold stress from that of the largest element-average stress in extension. The mean response is nonmonotonic for the uniform reference, so a monotonic Richardson extrapolation is not applied. Table 2 reports the change from 400 to 800 span elements. The same refinement halves all element dimensions, including through the gold thickness.')
figure('mesh_study','Figure 2. Tension mesh study. Means stabilize more rapidly than maxima. Each refinement also increases through-thickness resolution; the horizontal axis alone does not state total element count.')
mesh_data=[]
for n,a,b in zip(names,base,bend):
    old=get(n,r=8);oldb=get(n,'Bending',8)
    mesh_data.append([n,f"{100*abs(a['gold_mean_vm_MPa']-old['gold_mean_vm_MPa'])/a['gold_mean_vm_MPa']:.2f}",f"{100*abs(b['gold_mean_vm_MPa']-oldb['gold_mean_vm_MPa'])/b['gold_mean_vm_MPa']:.2f}",f"{100*abs(a['gold_max_cell_vm_MPa']-old['gold_max_cell_vm_MPa'])/a['gold_max_cell_vm_MPa']:.2f}"])
table('Table 2. Absolute mesh change relative to the 800-element result (%).',['Design','Mean tension','Mean transverse','Max. tension'],mesh_data)
para('The primary mean-stress changes are below 5%. This is a successive-mesh stability observation, not a 5% bound on the unknown continuum error. Gold mean axial strain and upper-tail stresses are less stable than mean equivalent stress. Maximum stresses are retained as diagnostic values and are not used to establish failure margins.')
h('4.2. Primary architecture comparison',True)
para(f'''Table 3 summarizes the finest-mesh extension response. Uniform PDMS gave {base[0]['gold_mean_vm_MPa']:.3f} MPa mean gold stress. The single buffer reduced this to {base[1]['gold_mean_vm_MPa']:.3f} MPa ({red[1]:.1f}% reduction), while grading reduced it to {base[2]['gold_mean_vm_MPa']:.3f} MPa ({red[2]:.1f}% reduction). The graded design therefore had {100*(base[2]['gold_mean_vm_MPa']/base[1]['gold_mean_vm_MPa']-1):.1f}% higher mean sensor stress than the single soft layer under this load. Figure 3 compares extension and transverse loading.''')
table('Table 3. Gold response under 1% nominal extension at 800 span elements. Stresses are in MPa.',['Design','Mean','95th percentile','Max. cell'],[[n,f"{a['gold_mean_vm_MPa']:.3f}",f"{a['gold_p95_vm_MPa']:.3f}",f"{a['gold_max_cell_vm_MPa']:.3f}"] for n,a in zip(names,base)])
figure('mean_stress','Figure 3. Gold mean equivalent stress for the two non-proportional load cases. Transverse motion is a prescribed skin-edge displacement, not pure bending.')
para('Figure 4 shows the gold-only stress fields on a common color scale. Larger values occur on the opposing gold surfaces within the sensor span rather than uniformly across the sensing region. Both compliant supports reduce the overall magnitude relative to the uniform reference, while preserving a pronounced spatial gradient. The plots show undeformed geometry, with the sensor thickness magnified by the panel aspect ratio.')
figure('gold_fields','Figure 4. Gold element-average equivalent stress in extension. Vertical coordinates are measured from the bottom of the gold layer. All panels use the same stress scale; no nodal stress averaging is applied.')
strain=[a['gold_mean_abs_exx']*1e6 for a in base]
para(f'''Gold mean absolute axial strains were {strain[0]:.2f}, {strain[1]:.2f}, and {strain[2]:.2f} microstrain for uniform, single-buffer, and graded supports. Relative to the imposed 10,000 microstrain nominal skin extension, the transfer ratios were {strain[0]/10000:.5f}, {strain[1]/10000:.5f}, and {strain[2]/10000:.5f}. The very small values reflect the localized stiff gold region and the selected indirect loading path; they must not be interpreted as a universal attenuation factor for an actual wearable device.''')
para(f'''In the transverse case, gold mean stress was {bend[0]['gold_mean_vm_MPa']:.3f}, {bend[1]['gold_mean_vm_MPa']:.3f}, and {bend[2]['gold_mean_vm_MPa']:.3f} MPa. The single and graded buffers reduced this measure by {100*(1-bend[1]['gold_mean_vm_MPa']/bend[0]['gold_mean_vm_MPa']):.1f}% and {100*(1-bend[2]['gold_mean_vm_MPa']/bend[0]['gold_mean_vm_MPa']):.1f}%, respectively. At the 200-span-element mesh, separately solved shortening cases reproduced one half of the extension equivalent stresses, consistent with linear load scaling. No independent compression ranking is inferred beyond that identity.''')
h('4.3. Compliance and material controls',True)
matchdiff=100*(cm[0]['gold_mean_vm_MPa']/base[2]['gold_mean_vm_MPa']-1);revdiff=100*(cm[1]['gold_mean_vm_MPa']/base[2]['gold_mean_vm_MPa']-1)
para(f'''Figure 5 compares the nominal buffers and the two controls. The compliance-matched homogeneous buffer gave {cm[0]['gold_mean_vm_MPa']:.3f} MPa mean gold stress, a {abs(matchdiff):.2f}% difference from the nominal graded support. Reversing the two graded layers gave {cm[1]['gold_mean_vm_MPa']:.3f} MPa, a {abs(revdiff):.2f}% difference. These effects are smaller than the single-to-graded separation and comparable to, or smaller than, the observed mean-stress mesh changes. They do not establish a robust benefit for a particular grading direction.''')
figure('controls','Figure 5. Tension controls at 800 span elements. Matched denotes a homogeneous buffer with E = 0.1667 MPa and the same elementary series shear compliance as the graded pair.')
sens=[]
for soft in [.05,.1,.2]:
    values=[]
    for name in ['Single','Graded']:
        a=get(name,r=8) if soft==.1 else next(a for a in controls if a['design']==name and a['refinement']==8 and a['soft_MPa']==soft and a['skin_MPa']==.2)
        values.append(a['gold_mean_vm_MPa'])
    sens.append([f'{soft:.2f}',*[f'{v:.3f}' for v in values]])
table('Table 4. Soft-layer modulus sensitivity in extension at 400 span elements. Gold mean stresses are in MPa.',['Soft E (MPa)','Single','Graded'],sens)
para('Within the tested modulus range, the single buffer retained a lower mean gold stress than the graded buffer. Increasing soft-layer stiffness increased transferred stress in both architectures. The skin-stiffness perturbations also retained the primary ordering: single buffer, graded buffer, then uniform PDMS in increasing mean stress. These are matched-mesh deterministic comparisons. The sweep demonstrates robustness to the stated perturbations, without establishing uncertainty bounds for fabricated devices.')
h('5. Discussion');paragraphs(text.DISCUSSION)
h('6. Limitations and next steps');paragraphs(text.LIMITS)
h('7. Conclusion')
para(f'''A reproducible plane-strain model of a layered wearable biosensor predicts lower mean gold stress with compliant support layers when total support thickness is controlled. Under 1% prescribed skin extension, mean sensor stress decreased by {red[1]:.1f}% with a single soft buffer and {red[2]:.1f}% with the specified graded pair relative to uniform PDMS. The single buffer also provided the lowest mean stress under the tested transverse displacement. A compliance-matched homogeneous buffer performed similarly to the graded support, indicating that overall compliance explains much of the observed isolation in this geometry.''')
para('The study supports evaluating grading against equal-thickness and compliance-matched baselines, with sensor-specific metrics and explicit mesh histories. It does not establish electrical stability, fatigue resistance, or safety. The next step is independent numerical comparison and experimental calibration of the materials, interfaces, and measured skin motion.')
h('Data and code availability')
para('The accompanying reproducibility package contains Python source, nominal and control metrics in CSV format, displacement and stress fields in compressed NumPy format, verification outputs, and source-provenance notes. Figure and manuscript generation are scripted. All numerical results reported here were generated by the specified elastic model; no human or animal measurements were collected.')
h('References')
for i,s in enumerate(text.REFERENCES,1):
    p=para(f'[{i}] {s}',10,indent=False);p.paragraph_format.left_indent=Cm(.55);p.paragraph_format.first_line_indent=Cm(-.55);p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.LEFT

for section in doc.sections:
    section.different_first_page_header_footer=False
    for hf in [section.header,section.first_page_header,section.even_page_header,section.footer,section.first_page_footer,section.even_page_footer]:
        for child in list(hf._element):hf._element.remove(child)
        p=OxmlElement('w:p');hf._element.append(p)
    p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
doc.core_properties.title=text.TITLE;doc.core_properties.author='Joel Whaling; Aarav Patel; Mandy Hung; Yuancheng Fang; Leona Cai; Zenan Li; Advik Raja; Nathan Xu; Linda Shi';doc.core_properties.subject='Reproducible finite-element study of compliant biosensor supports';doc.core_properties.comments=''
doc.save(OUT/'Biosensor_Proceeding.docx')
(OUT/'Manuscript_Text.txt').write_text('\n\n'.join(body_text),encoding='utf-8')
(OUT/'paper_statistics.json').write_text(json.dumps({'words':len(' '.join(body_text).split()),'figure_count':5,'table_count':4,'references':len(text.REFERENCES),'template_sha256':hashlib.sha256(template.read_bytes()).hexdigest(),'nominal_solved_cases':len(rows),'control_solved_cases':len(controls)},indent=2))
print((OUT/'paper_statistics.json').read_text())
